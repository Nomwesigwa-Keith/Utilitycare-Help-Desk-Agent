from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "Week2_Progress_Report.docx"
NAVY = "17365D"
LIGHT_BLUE = "D9EAF7"
GRAY = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), GRAY)
        borders.append(tag)
    tc_pr.append(borders)


def add_text(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.08
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
    paragraph.add_run(text)
    return paragraph


def add_heading(doc, text):
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor.from_string("000000")
    return paragraph


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    title_run = title.add_run("Week 2 Progress Report")
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle.add_run("UtiliCare Help Desk Triage Agent").italic = True

    add_text(
        doc,
        "Week 2 moved the project from planning artefacts toward a runnable triage baseline. "
        "The repository now contains a Gemini-backed FastAPI endpoint, a versioned prompt specification, "
        "and a repeatable ten-case evaluation harness. The evaluation infrastructure works and preserves real API outcomes; "
        "the configured Gemini project could not return classifications because its quota was exhausted or requests timed out.",
    )

    add_heading(doc, "Work completed")
    add_text(doc, "A FastAPI service exposes a health endpoint and a /triage endpoint. The endpoint accepts a customer message, calls the configured Gemini model, and returns both the raw model output and parsed JSON where available.")
    add_text(doc, "The active v1.1 prompt defines the assistant's role, six allowed categories, JSON-only response contract, safety limits, ambiguity handling, and fallback behavior. The change log retains v1.0 and explains why v1.1 added explicit safeguards and clarification fields.")
    add_text(doc, "Ten realistic evaluation cases now cover outages, billing, service requests, account access, complaints, an ambiguous multi-issue message, and nonsense input. The evaluator writes a CSV spreadsheet with expected category, returned category, pass/fail state, raw response, and parse or API errors.")

    add_heading(doc, "Decisions and evidence")
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.columns[0].width = Inches(1.25)
    table.columns[1].width = Inches(2.15)
    table.columns[2].width = Inches(3.55)
    headers = ["Contributor", "Evidence in repository", "Decision or contribution"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
            run.font.size = Pt(9)

    rows = [
        ("Keith Nomwesigwa", "Repository history and merge commits", "Set up the repository and integrated team branches, preserving GitHub as the project evidence trail."),
        ("Nankunda Lilian", "Project Charter commit 5e0c806", "Defined the problem, scope, constraints, and the requirement that routing, escalation, and closure remain human-approved."),
        ("Khot Adet Majuong", "Architecture document upload e603700", "Produced the initial architecture diagram, providing the system view for the agent and its tool boundaries."),
        ("Arok John Mayen", "Week 1 report commit a392929", "Recorded the Week 1 planning work, ownership, and unresolved category-taxonomy and model-selection questions."),
        ("Mubiru Umar", "User-story upload b383d3a and Gemini baseline commit 6dd6a34", "Defined acceptance criteria and integrated the initial Gemini-backed API baseline."),
    ]
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, text in enumerate(row):
            cells[index].text = text
            set_cell_border(cells[index])
            if row_index % 2 == 1:
                set_cell_shading(cells[index], LIGHT_BLUE)
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)

    add_heading(doc, "Problems encountered and response")
    add_text(doc, "The bundled Python runtime did not initially include the dependencies declared in requirements.txt, so the first evaluation could not import dotenv. The evaluation environment was completed in an isolated local runtime folder, and the evaluator was changed to capture per-case exceptions instead of stopping before it writes a results file.")
    add_text(doc, "The live Gemini evaluation recorded 0/10 passes. This is an authentic result, not a model-quality score: seven calls returned a 429 quota-exceeded response and three returned 504 deadline-exceeded responses. The result file retains those responses for audit. A valid API quota or an available model is required before classification quality can be measured.")

    add_heading(doc, "Week 3 priorities")
    add_text(doc, "Restore model access and rerun the ten cases, then expand evaluation toward the charter's target of 30 or more scenarios. Validate the category taxonomy against the synthetic knowledge base, add the mock outage lookup and source-grounded retrieval flow, and ensure a ticket remains a reviewable draft until an agent explicitly approves it.")

    add_heading(doc, "Repository evidence")
    add_text(doc, "Repository: https://github.com/Nomwesigwa-Keith/Utilitycare-Help-Desk-Agent/tree/utility-2. Evidence baseline reviewed for this report: commit a82ef81 (Add project documentation). The final submission commit should be supplied by the team member who performs the final push.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("UtiliCare Help Desk Triage Agent | Week 2")
    footer.runs[0].font.size = Pt(8)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
