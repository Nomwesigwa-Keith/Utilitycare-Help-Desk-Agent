from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "Model_Selection_Note.docx"
NAVY = "17365D"
PALE = "EAF2F8"
GRID = "D9D9D9"


def shade(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    properties.append(element)


def border(cell):
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), GRID)
        borders.append(element)
    properties.append(borders)


def paragraph(doc, text, size=9.5, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.0
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        run.font.size = Pt(size)
    run = p.add_run(text)
    run.font.size = Pt(size)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(9.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Model Selection Note")
    run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle.add_run("UtiliCare Help Desk Triage Agent").italic = True

    paragraph(doc, "Recommendation: use Gemini Flash for the Week 2 baseline, subject to restoring usable Gemini API quota. It fits the existing Python integration, is fast enough for short single-message classification, and supports JSON-formatted output. Gemini is not a final choice for real customer data: the project must continue to use synthetic inputs and confirm the selected billing and data-handling terms before any production use.")

    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    for column, width in zip(table.columns, [1.2, 1.7, 2.3, 1.8]):
        column.width = Inches(width)
    headings = ["Provider model", "Capability and latency", "Cost and access", "Privacy fit"]
    for index, text in enumerate(headings):
        cell = table.rows[0].cells[index]
        cell.text = text
        shade(cell, NAVY)
        border(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8)

    rows = [
        ("Google Gemini Flash\nRecommended", "Good fit for short classification and JSON output; Flash prioritizes low latency.", "Existing code and key path already use Gemini. Low entry cost, but the current free-tier quota blocked evaluation; budget or wait for quota before demo.", "Use synthetic messages only. Review the applicable Gemini API terms and data controls before sending any real customer data."),
        ("Anthropic Claude Haiku", "Fast Claude option with text, image, multilingual, and tool-use support. More capability than this baseline needs.", "Claude Haiku 4.5 is listed at $1 input and $5 output per million tokens. Requires an Anthropic account, key, and billing access.", "A viable alternative where Anthropic's account controls and terms meet the team's requirements; no key is currently configured."),
        ("OpenAI small model", "Suitable for structured classification and API integration; evaluate response quality and latency with the same test set.", "Usage-priced API with account access and an API key required. Compare actual token usage, not only list price, before choosing.", "API data controls are documented by OpenAI. The project's synthetic-data rule remains necessary regardless of provider."),
    ]
    for row_number, values in enumerate(rows):
        cells = table.add_row().cells
        for index, text in enumerate(values):
            cells[index].text = text
            border(cells[index])
            if row_number % 2 == 1:
                shade(cells[index], PALE)
            for p in cells[index].paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 0.95
                for run in p.runs:
                    run.font.size = Pt(7.8)

    paragraph(doc, "Decision rationale: Gemini Flash is selected for the baseline because it is already wired into the repository and is appropriate for the narrowly scoped triage classification task. The observed quota failures are an operational risk, not evidence that the model is inaccurate. Before the Week 3 demo, run the same ten cases after access is restored and keep Claude Haiku as the fallback if Gemini access remains unavailable.", 9)
    paragraph(doc, "Sources checked 11 September 2026: Anthropic model overview and pricing (platform.claude.com/docs); Google Gemini API pricing and terms (ai.google.dev/gemini-api); OpenAI API pricing and data controls (openai.com/api/pricing and platform.openai.com/docs/guides/your-data).", 7.5)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
